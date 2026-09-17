"""Source adapters for Step 6.

These adapters only expose rows and schema candidates. They do not know about
Neo4j, the AI services, or Supabase-specific production credentials.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable, Protocol


class RecordSource(Protocol):
    source_system: str
    source_object: str
    source_format: str

    def rows(self) -> list[dict[str, Any]]: ...


class RecordsSource:
    def __init__(
        self,
        rows: Iterable[dict[str, Any]],
        *,
        source_system: str,
        source_object: str,
        source_format: str = "records",
    ) -> None:
        self._rows = [dict(row) for row in rows]
        self.source_system = source_system
        self.source_object = source_object
        self.source_format = source_format

    def rows(self) -> list[dict[str, Any]]:
        return [dict(row) for row in self._rows]


class CSVSource:
    def __init__(
        self,
        path: str | Path,
        *,
        source_system: str = "csv",
        source_object: str | None = None,
        encoding: str = "utf-8-sig",
    ) -> None:
        self.path = Path(path)
        self.source_system = source_system
        self.source_object = source_object or self.path.name
        self.source_format = "csv"
        self.encoding = encoding

    def rows(self) -> list[dict[str, Any]]:
        with self.path.open("r", encoding=self.encoding, newline="") as handle:
            return [dict(row) for row in csv.DictReader(handle)]


class JSONArraySource:
    def __init__(
        self,
        path: str | Path,
        *,
        source_system: str = "json",
        source_object: str | None = None,
    ) -> None:
        self.path = Path(path)
        self.source_system = source_system
        self.source_object = source_object or self.path.name
        self.source_format = "json_array"

    def rows(self) -> list[dict[str, Any]]:
        payload = json.loads(self.path.read_text(encoding="utf-8"))
        if not isinstance(payload, list) or not all(isinstance(row, dict) for row in payload):
            raise ValueError("JSON source must contain a top-level array of objects")
        return [dict(row) for row in payload]


class XLSXSource:
    """Optional Excel adapter. openpyxl is imported lazily."""

    def __init__(
        self,
        path: str | Path,
        *,
        sheet_name: str | None = None,
        source_system: str = "xlsx",
        source_object: str | None = None,
    ) -> None:
        self.path = Path(path)
        self.sheet_name = sheet_name
        self.source_system = source_system
        self.source_object = source_object or (
            f"{self.path.name}#{sheet_name}" if sheet_name else self.path.name
        )
        self.source_format = "xlsx"

    def rows(self) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:  # pragma: no cover - environment-specific
            raise RuntimeError(
                "XLSX ingestion requires openpyxl. Install it with: pip install openpyxl"
            ) from error
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        sheet = workbook[self.sheet_name] if self.sheet_name else workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = next(values)
        except StopIteration:
            return []
        names = [str(item).strip() if item is not None else "" for item in headers]
        if any(not name for name in names):
            raise ValueError("Excel header row contains empty column names")
        rows: list[dict[str, Any]] = []
        for raw in values:
            rows.append({name: value for name, value in zip(names, raw)})
        return rows
